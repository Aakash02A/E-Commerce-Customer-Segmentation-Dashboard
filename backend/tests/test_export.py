"""Phase 10 export tests.

Cover CSV/Excel file contents, analysis isolation, missing/invalid/non-completed
analyses, empty cohorts, parameter validation, spreadsheet-formula-injection
mitigation and safe download headers.

Content tests reuse the real upload -> map -> validate -> aggregate -> RFM ->
segment pipeline (and the shared Spark session) so expectations are derived from
the same persisted results the dashboard reads -- never hardcoded numbers.
Edge cases that do not need Spark build analyses directly through the
repository.
"""
import csv
import io

from openpyxl import load_workbook

from app.repositories import analysis_repository
from conftest import prepare_ready_upload

STD_HEADER = "Client,Order_Ref,Purchase_Date,Units,Unit_Cost\n"


def _write(tmp_path, body, name="data.csv"):
    p = tmp_path / name
    p.write_text(STD_HEADER + body, encoding="utf-8")
    return str(p)


def _run_pipeline(client, csv_path):
    uid = prepare_ready_upload(client, csv_path)
    aid = client.post(f"/api/uploads/{uid}/aggregate", json={}).get_json()["data"]["analysis_id"]
    client.post(f"/api/analyses/{aid}/rfm", json={})
    client.post(f"/api/analyses/{aid}/segment", json={})
    return aid


def _export(client, aid, fmt="xlsx", dataset="full"):
    return client.get(f"/api/analyses/{aid}/export?format={fmt}&dataset={dataset}")


def _csv_rows(resp):
    text = resp.data.decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text)))


# --- Error handling -----------------------------------------------------
def test_export_unknown_analysis_404(client):
    assert _export(client, "doesnotexist", "xlsx", "full").status_code == 404
    assert _export(client, "doesnotexist", "csv", "customers").status_code == 404


def test_export_non_completed_analysis_409(client, tmp_path, spark):
    uid = prepare_ready_upload(client, _write(tmp_path, "A1,I1,2024-01-01,1,5.0\n"))
    with client.application.app_context():
        aid = analysis_repository.create_analysis(uid)  # pending, no metrics
    resp = _export(client, aid, "csv", "customers")
    assert resp.status_code == 409
    assert "completed" in resp.get_json()["message"].lower()


def test_export_invalid_format_400(client, tmp_path, spark):
    aid = _run_pipeline(client, _write(tmp_path, "A1,I1,2024-01-01,1,5.0\n"))
    assert _export(client, aid, fmt="pdf", dataset="customers").status_code == 400


def test_export_invalid_dataset_400(client, tmp_path, spark):
    aid = _run_pipeline(client, _write(tmp_path, "A1,I1,2024-01-01,1,5.0\n"))
    assert _export(client, aid, fmt="csv", dataset="bogus").status_code == 400


def test_export_csv_full_is_rejected_400(client, tmp_path, spark):
    aid = _run_pipeline(client, _write(tmp_path, "A1,I1,2024-01-01,1,5.0\n"))
    resp = _export(client, aid, fmt="csv", dataset="full")
    assert resp.status_code == 400


# --- Customers CSV: contents match persisted metrics --------------------
def test_customers_csv_matches_metrics(client, sample_csv_path, spark):
    aid = _run_pipeline(client, sample_csv_path)
    with client.application.app_context():
        metrics = [dict(m) for m in analysis_repository.get_metrics(aid)]

    resp = _export(client, aid, fmt="csv", dataset="customers")
    assert resp.status_code == 200
    rows = _csv_rows(resp)
    header = rows[0]
    assert header[:4] == ["Customer ID", "First Order Date", "Last Order Date", "Recency (days)"]
    data = rows[1:]
    assert len(data) == len(metrics)

    by_id = {r[0]: r for r in data}
    col = {name: i for i, name in enumerate(header)}
    for m in metrics:
        r = by_id[m["customer_id"]]
        assert r[col["Monetary (total spend)"]] == str(m["monetary"])
        assert r[col["Frequency (orders)"]] == str(m["frequency"])
        assert r[col["R Score"]] == str(m["r_score"])
        assert r[col["Segment"]] == (m["segment"] or "")


# --- Overview CSV: values match the dashboard overview ------------------
def test_overview_csv_matches_dashboard(client, sample_csv_path, spark):
    aid = _run_pipeline(client, sample_csv_path)
    dash = client.get(f"/api/analyses/{aid}/dashboard").get_json()["data"]
    ov = dash["overview"]

    rows = _csv_rows(_export(client, aid, fmt="csv", dataset="overview"))
    kv = {r[0]: r[1] for r in rows[1:]}
    assert kv["Total Customers"] == str(ov["total_customers"])
    assert kv["Total Valid Orders"] == str(ov["total_orders"])
    assert kv["Total Revenue"] == str(ov["total_revenue"])
    assert kv["Average Order Value"] == str(ov["average_order_value"])
    assert kv["AOV Definition"] == ov["aov_definition"]
    assert kv["Status"] == "completed"


# --- Segments CSV: counts/percentages match dashboard --------------------
def test_segments_csv_matches_dashboard(client, sample_csv_path, spark):
    aid = _run_pipeline(client, sample_csv_path)
    segs = client.get(f"/api/analyses/{aid}/dashboard").get_json()["data"]["segment_breakdown"]["segments"]

    rows = _csv_rows(_export(client, aid, fmt="csv", dataset="segments"))
    data = {r[0]: r for r in rows[1:]}
    assert len(data) == len(segs)
    for s in segs:
        assert data[s["segment"]][1] == str(s["count"])
        assert data[s["segment"]][2] == str(s["percentage"])


# --- Excel full workbook ------------------------------------------------
def test_xlsx_full_has_three_sheets(client, sample_csv_path, spark):
    aid = _run_pipeline(client, sample_csv_path)
    resp = _export(client, aid, fmt="xlsx", dataset="full")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Overview", "Customers", "Segments"]

    ws = wb["Customers"]
    all_rows = list(ws.iter_rows(values_only=True))
    with client.application.app_context():
        n = len(analysis_repository.get_metrics(aid))
    assert len(all_rows) - 1 == n  # header + one row per customer


def test_xlsx_customers_numeric_value(client, tmp_path, spark):
    aid = _run_pipeline(client, _write(tmp_path, "Solo,I1,2024-01-01,4,9.0\n"))
    resp = _export(client, aid, fmt="xlsx", dataset="customers")
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    monetary_idx = header.index("Monetary (total spend)")
    assert rows[1][monetary_idx] == 36.0  # 4 units * 9.0, numeric cell


# --- Analysis isolation -------------------------------------------------
def test_export_is_scoped_to_one_analysis(client, tmp_path, spark):
    a1 = _run_pipeline(client, _write(tmp_path, "A1,I1,2024-01-01,2,10.0\n", "ds1.csv"))
    a2 = _run_pipeline(client, _write(tmp_path, "B7,I9,2024-03-03,3,7.0\n", "ds2.csv"))

    ids1 = {r[0] for r in _csv_rows(_export(client, a1, "csv", "customers"))[1:]}
    ids2 = {r[0] for r in _csv_rows(_export(client, a2, "csv", "customers"))[1:]}
    assert ids1 == {"A1"} and ids2 == {"B7"}


# --- Formula-injection mitigation ---------------------------------------
def test_formula_injection_neutralized_in_csv(client, tmp_path, spark):
    evil = "=1+HYPERLEN"  # formula prefix, deliberately comma-free
    aid = _run_pipeline(
        client, _write(tmp_path, f"{evil},I1,2024-01-01,2,10.0\nNormal,I2,2024-01-02,1,5.0\n"))
    rows = _csv_rows(_export(client, aid, "csv", "customers"))
    ids = [r[0] for r in rows[1:]]
    # The dangerous id is exported as literal text (prefixed), never a live formula.
    assert f"'{evil}" in ids
    assert not any(v.startswith("=") for v in ids)


def test_formula_injection_neutralized_in_xlsx(client, tmp_path, spark):
    evil = "-1+1"
    aid = _run_pipeline(client, _write(tmp_path, f"{evil},I1,2024-01-01,2,10.0\n"))
    wb = load_workbook(io.BytesIO(_export(client, aid, "xlsx", "customers").data))
    cell = wb.active["A2"].value
    assert cell == f"'{evil}"


# --- Empty result handling ----------------------------------------------
def test_empty_cohort_export_is_headers_only(client, tmp_path, spark):
    uid = prepare_ready_upload(client, _write(tmp_path, "A1,I1,2024-01-01,1,5.0\n"))
    with client.application.app_context():
        aid = analysis_repository.create_analysis(uid)
        analysis_repository.record_completed(aid, {"transactions": {}})  # completed, no customers

    resp = _export(client, aid, "csv", "customers")
    assert resp.status_code == 200
    rows = _csv_rows(resp)
    assert len(rows) == 1  # header only, no data rows


# --- Download headers are safe ------------------------------------------
def test_download_headers_are_safe(client, tmp_path, spark):
    aid = _run_pipeline(client, _write(tmp_path, "A1,I1,2024-01-01,1,5.0\n"))
    resp = _export(client, aid, "csv", "customers")
    disp = resp.headers["Content-Disposition"]
    assert "attachment" in disp and "analysis_" in disp
    # No path separators or traversal in the filename.
    assert "/" not in disp and "\\" not in disp and ".." not in disp


def test_no_traceback_or_path_leakage_on_error(client):
    body = _export(client, "nope", "csv", "customers").get_data(as_text=True)
    assert "Traceback" not in body and "\\" not in body
