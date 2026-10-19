"""Validate uploaded files before reading them.

Checks extension, size, emptiness, magic-number/content sanity, and (for
Excel) lists worksheet names so the UI can offer a sheet selector. All checks
are non-destructive and never execute file contents or macros.
"""
import os
import zipfile

# Minimal magic numbers so we can detect a mismatched extension cheaply.
_XLSX_MAGIC = b"PK\x03\x04"  # ZIP container (Office Open XML)


class FileValidationError(Exception):
    """Raised with a user-facing message when a file cannot be accepted."""


def validate_file(file_storage, allowed_extensions: set[str], max_size: int) -> dict:
    """Validate a Werkzeug FileStorage object.

    Returns a dict: {extension, size, original_filename}. Raises
    FileValidationError with an understandable message on any failure.
    """
    if file_storage is None or not file_storage.filename:
        raise FileValidationError("No file was provided.")

    original_filename = file_storage.filename
    ext = (os.path.splitext(original_filename)[1] or "").lower().lstrip(".")

    if ext not in allowed_extensions:
        supported = ", ".join(sorted("." + e for e in allowed_extensions))
        raise FileValidationError(f"Unsupported file type '.{ext}'. Supported types: {supported}.")

    # Determine size reliably via seek/tell on the underlying stream.
    stream = file_storage.stream
    try:
        stream.seek(0, os.SEEK_END)
        size = stream.tell()
        stream.seek(0)
    except (OSError, AttributeError):
        size = 0

    if size == 0:
        raise FileValidationError("The uploaded file is empty.")
    if size > max_size:
        mb = max_size / (1024 * 1024)
        raise FileValidationError(f"File exceeds the maximum allowed size of {mb:.0f} MB.")

    _validate_content(ext, stream, size)
    stream.seek(0)

    return {
        "extension": ext,
        "size": size,
        "original_filename": original_filename,
    }


def _validate_content(ext: str, stream, size: int):
    """Light content sanity check that the bytes match the claimed type."""
    head = stream.read(8)
    stream.seek(0)
    if not head:
        raise FileValidationError("The file appears to be empty or unreadable.")

    if ext == "xlsx":
        if not head.startswith(_XLSX_MAGIC):
            raise FileValidationError(
                "This .xlsx file is not a valid Excel workbook (it may be corrupted, "
                "legacy .xls, or a mislabeled file)."
            )
    elif ext == "csv":
        # A CSV must start with readable text, not binary/control bytes.
        try:
            head.decode("utf-8")
        except UnicodeDecodeError:
            raise FileValidationError(
                "This .csv file does not appear to be plain text. Please export your "
                "data as CSV or Excel."
            )


def list_excel_sheets(path) -> list[str]:
    """Return worksheet names in an .xlsx file without loading all data."""
    if not zipfile.is_zipfile(path):
        raise FileValidationError("The Excel workbook could not be read (invalid or password-protected).")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    except Exception as exc:  # pragma: no cover - defensive
        raise FileValidationError(
            f"The Excel workbook could not be opened: {exc}. If it is password "
            "protected, remove protection and re-upload."
        )
